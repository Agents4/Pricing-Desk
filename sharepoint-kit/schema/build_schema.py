import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)

RIGHTS = [
    "internal_use", "internal_analytics_shared_outputs", "display_to_third_parties",
    "redistribution_resale", "derivative_data_product", "ai_training",
    "ai_fine_tuning", "ai_retrieval_grounding", "ai_evaluation",
]
RIGHT_STATUS = ["Granted", "Conditional", "Prohibited", "Unknown"]
YN = ["Yes", "No"]
SENS = ["Public", "Commercial-confidential", "Supplier-confidential", "Personal"]


def c(name, display, typ, required=False, choices=None, lookup=None, indexed=False, notes=""):
    return dict(name=name, display=display, type=typ, required=required, choices=choices or [],
                lookup=lookup or "", indexed=indexed, notes=notes)


TXT, MULTI, CHOICE, NUM, DATE, BOOL, PERSON, LOOKUP, LINK = (
    "Single line of text", "Multiple lines of text (plain)", "Choice", "Number", "Date (date only)",
    "Yes/No", "Person or Group", "Lookup", "Hyperlink",
)
LEGAL = [
    c("legal_confirmed", "Legal confirmed", CHOICE, choices=YN, indexed=True,
      notes="Written only by the approval flow. Hide from the new and edit forms. Blank counts as not confirmed."),
    c("confirmed_by", "Confirmed by", PERSON, notes="Written by the approval flow."),
    c("confirmed_date", "Confirmed date", DATE, notes="Written by the approval flow."),
]

LISTS = [
    dict(
        key="Parties", purpose="Every agency and supplier whose data or agreements affect rights.",
        key_column="party_id (the Title column, renamed)", approx_rows="Hundreds to low thousands",
        columns=[
            c("Title", "party_id", TXT, True, indexed=True, notes="Stable ID such as AGY-001 or AIR-A. Enforce unique values (index first). Rename the display name of Title."),
            c("name", "Name", TXT, True),
            c("party_type", "Party type", CHOICE, True, ["Supplier", "Agency"], indexed=True),
            c("region", "Region", TXT),
            c("status", "Status", CHOICE, True, ["Active", "Inactive"], indexed=True),
            c("relationship_owner", "Relationship owner", PERSON),
            c("notes", "Notes", MULTI),
        ],
    ),
    dict(
        key="Agreements", purpose="One row per contract, addendum, or amendment. The signed document sits in a locked library.",
        key_column="agreement_id (Title)", approx_rows="Parties times amendments, low thousands",
        columns=[
            c("Title", "agreement_id", TXT, True, indexed=True, notes="Unique, such as AGR-101."),
            c("party", "Party", LOOKUP, True, lookup="Parties > Title", indexed=True),
            c("profile", "Rights profile", LOOKUP, True, lookup="Rights_Profiles > Title", indexed=True, notes="Which standard terms this agreement follows."),
            c("agreement_type", "Agreement type", CHOICE, True, ["Content distribution", "Distribution", "Private fare addendum", "API terms", "Data license", "Amendment", "Other"]),
            c("status", "Status", CHOICE, True, ["Draft", "Active", "Superseded", "Terminated", "Expired"], indexed=True),
            c("effective_date", "Effective date", DATE, True),
            c("expiry_date", "Expiry date", DATE, indexed=True, notes="Blank means open-ended. The resolver fails closed after this date."),
            c("auto_renews", "Auto-renews", BOOL),
            c("supersedes", "Supersedes", LOOKUP, lookup="Agreements > Title", notes="Set on an amendment or renewal. A flow marks the earlier agreement Superseded."),
            c("deviates_from_profile", "Deviates from profile", BOOL, notes="Set by a flow when any exception exists."),
            c("document_link", "Signed document", LINK, True, notes="Link into the locked Contracts library."),
        ] + LEGAL + [c("notes", "Notes", MULTI)],
    ),
    dict(
        key="Rights_Profiles", purpose="Named sets of standard terms, such as Standard Airline v2. Confirmed once, reused by many agreements.",
        key_column="profile_id (Title)", approx_rows="Tens",
        columns=[
            c("Title", "profile_id", TXT, True, indexed=True, notes="Example: PROF-AIR-STD-2."),
            c("name", "Name", TXT, True),
            c("applies_to", "Applies to", CHOICE, True, ["Supplier", "Agency"]),
            c("version", "Version", NUM, True),
            c("status", "Status", CHOICE, True, ["Draft", "Approved", "Retired"], indexed=True, notes="Only Approved profiles count. Set by the approval flow."),
            c("approved_by", "Approved by", PERSON),
            c("approved_date", "Approved date", DATE),
            c("next_review", "Next review", DATE),
        ],
    ),
    dict(
        key="Profile_Terms", purpose="The nine use rights for each profile. One row per profile and use right.",
        key_column="profile + use_right", approx_rows="Profiles times 9",
        columns=[
            c("Title", "term_key", TXT, True, indexed=True, notes="Set by a flow to profile_id|use_right. Enforce unique values."),
            c("profile", "Profile", LOOKUP, True, lookup="Rights_Profiles > Title", indexed=True),
            c("profile_id", "profile_id", TXT, True, indexed=True, notes="Plain text copy of the profile ID, set by a flow, so exports are self-contained."),
            c("use_right", "Use right", CHOICE, True, RIGHTS, indexed=True),
            c("status", "Status", CHOICE, True, RIGHT_STATUS, indexed=True),
            c("conditions", "Conditions", MULTI, notes="Conditions and limits as stated in the evidence."),
            c("evidence_ref", "Evidence", LOOKUP, lookup="Evidence > Title", notes="Allow multiple values."),
        ] + LEGAL,
    ),
    dict(
        key="Agreement_Exceptions", purpose="Only the deviations from a party's profile. An exception overrides the profile term for one right, and can loosen or tighten it.",
        key_column="exception_id (Title)", approx_rows="Small. Hundreds at most",
        columns=[
            c("Title", "exception_id", TXT, True, indexed=True, notes="Example: EX-001."),
            c("agreement", "Agreement", LOOKUP, True, lookup="Agreements > Title", indexed=True),
            c("agreement_id", "agreement_id", TXT, True, indexed=True, notes="Plain text copy set by a flow, for exports."),
            c("dataset", "Dataset", LOOKUP, lookup="Datasets > Title", notes="Blank means all datasets."),
            c("dataset_id", "dataset_id", TXT, notes="Plain text copy set by a flow. Blank means all datasets."),
            c("use_right", "Use right", CHOICE, True, RIGHTS, indexed=True),
            c("status", "Status", CHOICE, True, RIGHT_STATUS),
            c("clause_ref", "Clause reference", TXT, True),
            c("reason", "Reason", MULTI, True, notes="Turn on append-only version history so earlier text is kept."),
            c("effective_date", "Effective date", DATE),
            c("expiry_date", "Expiry date", DATE),
            c("evidence_ref", "Evidence", LOOKUP, lookup="Evidence > Title"),
        ] + LEGAL,
    ),
    dict(
        key="Datasets", purpose="The logical datasets. As in the first register, plus the derivation fields the resolver needs.",
        key_column="dataset_id (Title)", approx_rows="Tens to low hundreds",
        columns=[
            c("Title", "dataset_id", TXT, True, indexed=True, notes="DS-<SOURCE>-<NNN>. Never reused."),
            c("name", "Name", TXT, True),
            c("description", "Description", MULTI, True),
            c("owner_role", "Owner role", TXT, True),
            c("source_type", "Source type", CHOICE, True, ["Owned", "Supplier-sourced", "Customer-contributed", "Third-party licensed", "Public", "Derived"]),
            c("contributor_party_type", "Contributor party type", CHOICE, choices=["Supplier", "Agency"], notes="For non-derived datasets: which kind of party contributes the data."),
            c("parent_dataset_ids", "Parent dataset IDs", TXT, notes="Semicolon separated. Plain text so exports match the resolver."),
            c("derivation_mode", "Derivation mode", CHOICE, choices=["Inherit", "Transform"],
              notes="Inherit for joins and copies. Transform for aggregated or anonymized products: parents decide who is included, the dataset's own ceiling sets output rights."),
            c("coverage", "Coverage", MULTI),
            c("refresh_frequency", "Refresh frequency", TXT),
            c("history_depth", "History depth", TXT),
            c("delivery_methods", "Delivery methods", TXT),
            c("contains_personal_data", "Contains personal data", CHOICE, choices=["Yes", "No", "Unknown"]),
            c("sensitivity", "Sensitivity", CHOICE, choices=SENS),
            c("status", "Status", CHOICE, True, ["Draft", "In legal review", "Approved", "Restricted", "Rejected", "Retired"], indexed=True),
            c("version", "Version", NUM),
            c("next_review", "Next review", DATE),
            c("approved_by", "Approved by", PERSON),
            c("approved_date", "Approved date", DATE),
        ],
    ),
    dict(
        key="Dataset_Use_Rights", purpose="The dataset-level ceiling for each use right: what the company may ever do with the dataset as a whole, whatever any one contract says. Set once per dataset and right.",
        key_column="dataset + use_right", approx_rows="Datasets times 9",
        columns=[
            c("Title", "ceiling_key", TXT, True, indexed=True, notes="Set by a flow to dataset_id|use_right. Enforce unique values."),
            c("dataset", "Dataset", LOOKUP, True, lookup="Datasets > Title", indexed=True),
            c("dataset_id", "dataset_id", TXT, True, indexed=True, notes="Plain text copy set by a flow."),
            c("use_right", "Use right", CHOICE, True, RIGHTS, indexed=True),
            c("status", "Status", CHOICE, True, RIGHT_STATUS),
            c("conditions", "Conditions", MULTI),
            c("evidence_ref", "Evidence", LOOKUP, lookup="Evidence > Title"),
        ] + LEGAL,
    ),
    dict(
        key="Evidence", purpose="Pointers and short plain-language summaries. Contract text stays in the locked library.",
        key_column="evidence_id (Title)", approx_rows="Thousands",
        columns=[
            c("Title", "evidence_id", TXT, True, indexed=True, notes="EV-NNNN."),
            c("agreement", "Agreement", LOOKUP, lookup="Agreements > Title", indexed=True),
            c("dataset", "Dataset", LOOKUP, lookup="Datasets > Title"),
            c("document_type", "Document type", CHOICE, True, ["Supplier agreement", "Customer agreement", "Terms of service", "Policy", "Email", "Other"]),
            c("location", "Location", TXT, True, notes="Section, clause, or page."),
            c("summary", "Summary", MULTI, True, notes="One or two sentences in your own words. Do not paste clause text."),
            c("document_link", "Document", LINK),
        ],
    ),
    dict(
        key="Fields", purpose="Field-level sensitivity and restricted uses. Unchanged from the first register, plus a dataset lookup.",
        key_column="dataset + field_name", approx_rows="Hundreds to thousands",
        columns=[
            c("Title", "field_key", TXT, True, indexed=True, notes="dataset_id|field_name."),
            c("dataset", "Dataset", LOOKUP, True, lookup="Datasets > Title", indexed=True),
            c("field_name", "Field name", TXT, True),
            c("description", "Description", MULTI),
            c("data_type", "Data type", CHOICE, choices=["String", "Number", "Date", "Boolean", "JSON", "Other"]),
            c("sensitivity", "Sensitivity", CHOICE, True, SENS),
            c("origin", "Origin", TXT),
            c("restricted_uses", "Restricted uses", TXT),
            c("notes", "Notes", MULTI),
        ],
    ),
    dict(
        key="Open_Questions", purpose="Questions for Legal, Product, Data engineering, and Finance. Now optionally tied to a party or agreement.",
        key_column="question_id (Title)", approx_rows="Hundreds",
        columns=[
            c("Title", "question_id", TXT, True, indexed=True, notes="Q-NNNN."),
            c("dataset", "Dataset", LOOKUP, lookup="Datasets > Title", indexed=True),
            c("party", "Party", LOOKUP, lookup="Parties > Title"),
            c("agreement", "Agreement", LOOKUP, lookup="Agreements > Title"),
            c("question", "Question", MULTI, True),
            c("owner_role", "Owner role", CHOICE, True, ["Legal", "Product", "Data engineering", "Finance", "Sales operations"]),
            c("priority", "Priority", CHOICE, True, ["High", "Medium", "Low"]),
            c("status", "Status", CHOICE, True, ["Open", "Answered", "Closed"], indexed=True),
            c("due_date", "Due date", DATE),
            c("resolution", "Resolution", MULTI),
        ],
    ),
    dict(
        key="Approvals_Log", purpose="Written only by the approval flow. Records every Legal decision, so a confirmed flag can be traced to an approval.",
        key_column="approval_id (Title)", approx_rows="Grows with every decision",
        columns=[
            c("Title", "approval_id", TXT, True, indexed=True, notes="Flow run ID."),
            c("item_list", "List", TXT, True, indexed=True),
            c("item_id", "Item ID", NUM, True, indexed=True),
            c("item_version", "Item version", TXT),
            c("decision", "Decision", CHOICE, True, ["Approved", "Rejected"]),
            c("approver", "Approver", PERSON, True),
            c("decided_on", "Decided on", DATE, True),
            c("comment", "Comment", MULTI),
        ],
    ),
    dict(
        key="Effective_Rights", purpose="Generated by the resolver nightly. Read-only. Move to the warehouse if it passes about 100,000 rows.",
        key_column="dataset_id + party_id + use_right", approx_rows="Datasets times parties times 9",
        columns=[
            c("Title", "row_key", TXT, True, indexed=True, notes="dataset_id|party_id|use_right."),
            c("dataset_id", "dataset_id", TXT, True, indexed=True),
            c("party_id", "party_id", TXT, True, indexed=True),
            c("party_type", "Party type", CHOICE, choices=["Supplier", "Agency"]),
            c("use_right", "Use right", CHOICE, True, RIGHTS, indexed=True),
            c("effective_status", "Effective status", CHOICE, True, RIGHT_STATUS),
            c("eligible", "Eligible", CHOICE, True, YN, indexed=True),
            c("basis", "Basis", TXT),
            c("agreement_id", "Agreement", TXT),
            c("as_of_date", "As of", DATE, True),
        ],
    ),
    dict(
        key="Coverage_Summary", purpose="Generated by the resolver nightly. Share of parties and volume that may be used for each right.",
        key_column="dataset_id + use_right + basis", approx_rows="Datasets times 9 times a few",
        columns=[
            c("Title", "row_key", TXT, True, indexed=True),
            c("dataset_id", "dataset_id", TXT, True, indexed=True),
            c("use_right", "Use right", CHOICE, True, RIGHTS),
            c("basis", "Basis", TXT, True, notes="Agency parties, Supplier parties, or Joint agency and supplier volume."),
            c("total", "Total", NUM),
            c("eligible", "Eligible", NUM),
            c("pct_by_party", "Percent by party", NUM),
            c("pct_by_volume", "Percent by volume", NUM),
            c("as_of_date", "As of", DATE),
        ],
    ),
]

# group -> list -> access
GROUPS = ["Pricing readers", "Register stewards", "Legal approvers", "Resolver service account"]
ACCESS = {
    "Parties": ["Read", "Contribute", "Approve", "Read"],
    "Agreements": ["No access", "Contribute", "Approve", "Read"],
    "Contracts library (signed documents)": ["No access", "Contribute", "Approve", "No access"],
    "Rights_Profiles": ["Read", "Contribute", "Approve", "Read"],
    "Profile_Terms": ["Read", "Contribute", "Approve", "Read"],
    "Agreement_Exceptions": ["No access", "Contribute", "Approve", "Read"],
    "Datasets": ["Read", "Contribute", "Approve", "Read"],
    "Dataset_Use_Rights": ["Read", "Contribute", "Approve", "Read"],
    "Evidence": ["No access", "Contribute", "Approve", "No access"],
    "Fields": ["Read", "Contribute", "Approve", "No access"],
    "Open_Questions": ["Read", "Contribute", "Contribute", "No access"],
    "Approvals_Log": ["No access", "Read", "Read", "Contribute (flow only)"],
    "Effective_Rights": ["Read", "Read", "Read", "Contribute"],
    "Coverage_Summary": ["Read", "Read", "Read", "Contribute"],
}

FLOWS = [
    ("F1 Legal approval", "An item in Profile_Terms, Agreement_Exceptions, Dataset_Use_Rights, Agreements, or Rights_Profiles is submitted for approval",
     "Start an approval for the Legal group. On approve: set legal_confirmed = Yes, confirmed_by, confirmed_date (or profile status = Approved). On reject: set legal_confirmed = No and post the comment. Always write a row to Approvals_Log.",
     "Approvals_Log, legal_confirmed fields"),
    ("F2 Exception guard", "Agreement_Exceptions item created or changed",
     "Set Agreements.deviates_from_profile = Yes. Clear legal_confirmed on the exception. Start F1. Copy agreement_id and dataset_id into the text columns.",
     "Agreements, Agreement_Exceptions"),
    ("F3 Amendment", "Agreements item set to Active with a Supersedes value",
     "Set the superseded agreement to Superseded. Notify the register steward and Legal.",
     "Agreements"),
    ("F4 Expiry watch", "Daily",
     "Find Active agreements expiring in 90, 60, and 30 days and notify the steward and Legal. On the expiry date, set status = Expired. Rights then resolve to Unknown.",
     "Agreements"),
    ("F5 New party intake", "Parties item created",
     "Create a task for the steward to add the agreement, pick a profile, and upload the document. The party stays excluded (no active confirmed agreement) until Legal approves.",
     "Tasks"),
    ("F6 Review dates", "Daily",
     "Notify owners 30 days before next_review on Datasets and Rights_Profiles. Reopen to In legal review when a source, schema, contract, or use-right request changes.",
     "Datasets, Rights_Profiles"),
    ("F7 Key columns", "Item created in Profile_Terms, Dataset_Use_Rights, Fields",
     "Fill Title and the plain-text ID columns from the lookups, so every export is self-contained.",
     "Title and ID columns"),
    ("F8 Nightly resolver", "Daily, after F4",
     "Export the lists to CSV (Microsoft Graph or the list export), run resolve_effective_rights.py, write effective_rights, coverage_summary, and issues to Effective_Rights, Coverage_Summary, and the warehouse. Notify Legal and the steward of new Unknown or Prohibited results and warnings.",
     "Effective_Rights, Coverage_Summary"),
]

VIEWS = [
    ("Agreements", "Expiring in 90 days", "status = Active and expiry_date within 90 days. Sort by expiry_date."),
    ("Agreements", "Awaiting Legal", "status = Active or Draft and legal_confirmed is not Yes."),
    ("Agreement_Exceptions", "By party", "Group by agreement. Show use_right, status, dataset, clause_ref."),
    ("Profile_Terms", "Matrix by profile", "Group by profile. One column per use right is easier in Power BI."),
    ("Dataset_Use_Rights", "Unconfirmed ceilings", "legal_confirmed is not Yes."),
    ("Effective_Rights", "Eligible by dataset and right", "Filter eligible = Yes. Group by dataset_id then use_right."),
    ("Coverage_Summary", "Coverage by dataset", "Group by dataset_id. Show pct_by_party and pct_by_volume."),
    ("Open_Questions", "Open for Legal", "owner_role = Legal and status = Open. Sort by priority, due_date."),
]

NOTES = [
    "Resolution rule, per dataset, party, and use right:",
    "1. Find the party's Active, confirmed agreement in force on the as-of date. If none, the result is Unknown (fail closed).",
    "2. The agreement's profile must be Approved. The profile term must be confirmed, or the result is Unknown.",
    "3. An exception on the agreement (for this dataset or all datasets, in date range) replaces the profile term. It can loosen or tighten. An unconfirmed exception makes the term Unknown.",
    "4. The dataset-level ceiling (Dataset_Use_Rights) caps the contract term. A missing or unconfirmed ceiling is Unknown.",
    "5. Derived datasets: Inherit takes the most restrictive result across parents for the same right. Transform (aggregated or anonymized) uses the parents only to decide whether a party's data may be included (the derivative_data_product right), and takes output rights from the dataset's own ceiling.",
    "Ranking, most restrictive first: Prohibited, Unknown, Conditional, Granted. Eligible means Granted or Conditional.",
    "",
    "Build notes:",
    "Use the built-in Title column as each list's ID and rename its display name. Index a column before turning on Enforce unique values.",
    "Internal names matter: create each column with the name in the Column sheets (no spaces) and rename the display name afterward, so exports match the resolver's CSV headers.",
    "Plain-text ID copies (profile_id, agreement_id, dataset_id) sit beside lookups because lookup exports show display values, and the resolver needs stable IDs.",
    "A single view can show at most 12 lookup-type columns, and list views over 5,000 items need indexed filter columns. Index every column marked Indexed.",
    "SharePoint cannot restrict one column to one group. Hide legal_confirmed fields on the forms and let only the approval flow write them. The Approvals_Log is the control: any confirmed row without a matching log entry should be investigated.",
    "Turn on versioning (major versions) and content approval on every authored list. Keep version history for the audit trail.",
    "Give each list unique permissions. Pricing never gets access to Agreements, Agreement_Exceptions, Evidence, or the Contracts library.",
    "Sample CSV headers for the resolver are in sample_data. Each export must keep those column names.",
]

# ---------------------------------------------------------------- workbook
wb = Workbook()
ws = wb.active
ws.title = "README"
ws["A1"] = "SharePoint list schema: party-aware data rights register"
ws["A1"].font = Font(bold=True, size=13)
for i, line in enumerate(NOTES, start=3):
    ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

hdr_fill = PatternFill("solid", fgColor="E8EEF5")


def header(sheet, titles, widths):
    sheet.append(titles)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    sheet.freeze_panes = "A2"


ls = wb.create_sheet("Lists")
header(ls, ["List", "Purpose", "Key", "Approximate size"], [26, 80, 38, 34])
for L in LISTS:
    ls.append([L["key"], L["purpose"], L["key_column"], L["approx_rows"]])
for row in ls.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

for L in LISTS:
    sh = wb.create_sheet(L["key"][:31])
    header(sh, ["Internal name", "Display name", "Type", "Required", "Choices", "Lookup to", "Indexed", "Notes"],
           [24, 24, 28, 10, 50, 28, 9, 70])
    for col in L["columns"]:
        sh.append([col["name"], col["display"], col["type"], "Yes" if col["required"] else "", "; ".join(col["choices"]),
                   col["lookup"], "Yes" if col["indexed"] else "", col["notes"]])
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

pm = wb.create_sheet("Permissions")
header(pm, ["List or library"] + GROUPS, [40, 20, 20, 20, 26])
for k, v in ACCESS.items():
    pm.append([k] + v)
pm.append([])
pm.append(["Approve = SharePoint Approve permission with content approval on. Contribute = add and edit items. Read = view approved items."])

fl = wb.create_sheet("Flows")
header(fl, ["Flow", "Trigger", "Steps", "Writes to"], [24, 50, 110, 34])
for f in FLOWS:
    fl.append(list(f))
for row in fl.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

vw = wb.create_sheet("Views")
header(vw, ["List", "View", "Definition"], [26, 34, 100])
for v in VIEWS:
    vw.append(list(v))

wb.save(OUT / "sharepoint-rights-register-schema.xlsx")

spec = dict(
    version=1,
    resolution_notes=NOTES,
    lists=[{k: L[k] for k in ("key", "purpose", "key_column", "approx_rows", "columns")} for L in LISTS],
    permissions=dict(groups=GROUPS, access=ACCESS),
    flows=[dict(name=f[0], trigger=f[1], steps=f[2], writes=f[3]) for f in FLOWS],
    views=[dict(list=v[0], view=v[1], definition=v[2]) for v in VIEWS],
)
(OUT / "sharepoint-rights-register-schema.json").write_text(json.dumps(spec, indent=2), encoding="utf-8")
print("lists:", len(LISTS), "columns:", sum(len(L["columns"]) for L in LISTS))
